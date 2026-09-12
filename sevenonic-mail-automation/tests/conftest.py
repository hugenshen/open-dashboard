from __future__ import annotations

import io
from email.message import EmailMessage
from pathlib import Path

import pytest
from openpyxl import Workbook

from sevenonic_mail.config import MailboxConfig, Settings


def build_tech_pack(*, revision: int = 4, bust_value: float = 36.0) -> bytes:
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "3. MEASUREMENT"
    measurement["A1"] = "HELLOMOLLY"
    measurement["A2"] = "HMV071"
    measurement.append([])
    measurement.append([])
    measurement.append(
        [
            "REF #",
            "MEASUREMENT POINT",
            "TOL +/-",
            "1ST GIVEN",
            "1ST FIT MEASURE",
            "1ST FIT COMMENT",
            "UPDATED FOR 2ND FIT",
        ]
    )
    measurement.append([1, "FRONT LENGTH FROM HSP TO HEM", 1.0, 150.0, 148.0, "increase", 150.0])
    measurement.append(
        [2, "1/2 BUST CIRC. 2.5CM BELOW U/ARM", 0.5, 35.0, 34.5, "revise", bust_value]
    )

    graded = workbook.create_sheet("4. Graded Specs (2)")
    graded.append(["REF #", "MEASUREMENT POINT", "TOL +/-", "US 4 / AU 8"])
    graded.append([1, "FRONT LENGTH FROM HSP TO HEM", 1.0, 150.0])
    graded.append([2, "1/2 BUST CIRC. 2.5CM BELOW U/ARM", 0.5, bust_value])

    comments = workbook.create_sheet("6. PP COMMENTS")
    comments["A1"] = "PP COMMENTS"
    comments["A2"] = "Front hem accepted; revise bust only"

    workbook.create_sheet("7. PHOTO SAMPLE APPROVAL")["A1"] = "PENDING APPROVAL"
    full_size = workbook.create_sheet("8. FULL SIZE MEASURE BULK")
    full_size.append(["REF #", "MEASUREMENT POINT", "US 2", "US 4", "US 6"])
    full_size.append([1, "STRAP LENGTH", -1.0, 0.0, 1.0])
    metadata = workbook.create_sheet("Metadata")
    metadata["A1"] = f"Revision {revision}"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_mime(workbook_bytes: bytes) -> bytes:
    message = EmailMessage()
    message["From"] = "buyer@customer-a.com"
    message["To"] = "po@sevenonic.com"
    message["Cc"] = "merch.a@sevenonic.com, boss@sevenonic.com"
    message["Subject"] = "HMV071 Tech Pack Rev 4"
    message["Message-ID"] = "<hmv071-r4@customer-a.com>"
    message.set_content("Please review attached HMV071 revision 4 tech pack.")
    message.add_alternative(
        "<html><body><p>Please review <b>HMV071 Rev 4</b>.</p></body></html>",
        subtype="html",
    )
    message.add_attachment(
        workbook_bytes,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="HMV071-Tech-Pack-Rev4.xlsx",
    )
    return message.as_bytes()


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        app_env="test",
        database_url=f"sqlite:///{tmp_path / 'test.db'}",
        artifact_root=str(tmp_path / "artifacts"),
        admin_token="test-admin-token",
        graph_client_state="test-client-state",
        graph_notification_url="https://example.test/webhooks/graph",
        graph_lifecycle_url="https://example.test/webhooks/graph/lifecycle",
        graph_mailboxes_json=[MailboxConfig(mailbox_id="po@sevenonic.com", folder_id="inbox")],
        monday_board_id="18430298578",
        monday_group_id="ai_review",
    )
